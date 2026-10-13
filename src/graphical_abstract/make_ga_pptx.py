"""Editable graphical abstract (PPTX): native shapes, text boxes, connectors and pictures, grouped by block."""
import sys as _sys
from pathlib import Path as _P

_sys.path.insert(0, str(_P(__file__).resolve().parents[2]))
from s2sr.paths import S2V_NPZ, S2V_SPLITS, WEIGHTS, WORK_DIR  # noqa: E402,F401
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = WORK_DIR
A = WORK_DIR / "ga_assets"
STEM = "mask84_039_075_15_18"
OUT = WORK_DIR / "graphical_abstract.pptx"

PX = 914400 / 120                      # 1 canvas pixel = 1/120 inch
W_PX, H_PX = 1600, 1330

NAVY = RGBColor(0x14, 0x21, 0x3D)
BLUE_BG, BLUE_EDGE = RGBColor(0xE3, 0xEC, 0xF8), RGBColor(0x6F, 0x8F, 0xC0)
GREEN_BG, GREEN_EDGE = RGBColor(0xE5, 0xF3, 0xE6), RGBColor(0x79, 0xA9, 0x7B)
YELLOW_BG, YELLOW_EDGE = RGBColor(0xFF, 0xF1, 0xB8), RGBColor(0xD4, 0xB6, 0x4A)
SEG_BG, SEG_EDGE = RGBColor(0xA9, 0xE6, 0xB0), RGBColor(0x4F, 0x9B, 0x5A)
CAT_BG, CAT_EDGE = RGBColor(0xB9, 0xD3, 0xF5), RGBColor(0x4A, 0x6F, 0xA8)
GREY_BG, GREY_EDGE = RGBColor(0xFA, 0xFA, 0xFA), RGBColor(0x9A, 0x9A, 0x9A)
MASK_BG = RGBColor(0xF4, 0xF4, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CLEAR, CLOUD, SHADOW = RGBColor(0xED, 0xED, 0xED), RGBColor(0x33, 0x8C, 0xF2), RGBColor(0x73, 0x26, 0x8C)
FONT = "Arial"
FS_T, FS_N, FS_S = 17, 16, 14


def E(v):
    return Emu(int(round(v * PX)))


prs = Presentation()
prs.slide_width = E(W_PX)
prs.slide_height = E(H_PX)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def rbox(shapes, x, y, w, h, fill, edge, r=18, lw=1.75, name=None):
    sh = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    sh.adjustments[0] = min(0.5, r / min(w, h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = edge
    sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if name:
        sh.name = name
    return sh


def label(shapes, cx, cy, w, h, text, size=FS_T, bold=True, color=NAVY, align=PP_ALIGN.CENTER, name=None, x_left=None):
    x = x_left if x_left is not None else cx - w / 2
    tb = shapes.add_textbox(E(x), E(cy - h / 2), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    if name:
        tb.name = name
    return tb


def pic(shapes, path, x, y, w, h, name=None, border=True):
    p = shapes.add_picture(str(path), E(x), E(y), E(w), E(h))
    if border:
        p.line.color.rgb = NAVY
        p.line.width = Pt(1.25)
    if name:
        p.name = name
    return p


def seg(shapes, x1, y1, x2, y2, arrow=False, color=NAVY, lw=2.25, name=None):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "lg")
        tail.set("len", "lg")
    if name:
        c.name = name
    return c


def group(name):
    g = slide.shapes.add_group_shape()
    g.name = name
    return g.shapes


# --- frame
frame = rbox(slide.shapes, 6, 6, W_PX - 12, H_PX - 12, GREY_BG, NAVY, r=26, lw=2.25, name="Frame")

# --- inputs
g = group("Block: Sentinel-2 patch")
rbox(g, 34, 34, 330, 410, BLUE_BG, BLUE_EDGE, name="Box")
label(g, 199, 90, 300, 100, "Sentinel-2\npatches\n10 m, 8 bands", name="Title")
pic(g, A / f"{STEM}_input_rgb.png", 74, 160, 250, 250, name="Picture: Sentinel-2 RGB")

g = group("Block: Sen2VENuS")
rbox(g, 402, 34, 330, 410, BLUE_BG, BLUE_EDGE, name="Box")
label(g, 567, 90, 300, 100, "Sen2VENµS\n5 m reference\nSR training", name="Title")
pic(g, A / "s2v_reference.png", 442, 190, 250, 215, name="Picture: VENuS 5 m reference")
label(g, 567, 425, 250, 30, "reference", FS_S, bold=False, name="Caption")

seg(slide.shapes, 199, 444, 199, 472, name="Line 1")
seg(slide.shapes, 567, 444, 567, 472, name="Line 2")
seg(slide.shapes, 199, 472, 567, 472, name="Line 3")
seg(slide.shapes, 383, 472, 383, 490, arrow=True, name="Arrow to SR")

# --- SR bar
g = group("Block: Upsampling bar")
rbox(g, 34, 490, 698, 104, YELLOW_BG, YELLOW_EDGE, name="Box")
label(g, 383, 542, 660, 80, "Upsampling to the 5 m grid\nEDiffSR · ESRGAN · bicubic", name="Title")

# --- SR outputs
cols = [(34, "Bicubic\ninterpolation", f"{STEM}_bicubic_rgb.png", "Bicubic"),
        (270, "EDiffSR\n5 m", f"{STEM}_ediffsr_rgb.png", "EDiffSR"),
        (506, "ESRGAN\n5 m", f"{STEM}_esrgan_rgb.png", "ESRGAN")]
for x, txt, f, nm in cols:
    seg(slide.shapes, x + 113, 594, x + 113, 660, arrow=True, name=f"Arrow to {nm}")
    g = group(f"Block: {nm} output")
    rbox(g, x, 660, 226, 400, GREEN_BG, GREEN_EDGE, name="Box")
    label(g, x + 113, 712, 200, 70, txt, FS_T, name="Title")
    pic(g, A / f, x + 18, 800, 190, 190, name=f"Picture: {nm}")
    seg(slide.shapes, x + 113, 1060, x + 113, 1088, name=f"Line {nm}")
seg(slide.shapes, 147, 1088, 619, 1088, name="Collector")
seg(slide.shapes, 619, 1088, 776, 1088, arrow=True, name="Arrow to concatenation")

# --- concatenation bar
cat = rbox(slide.shapes, 776, 440, 72, 650, CAT_BG, CAT_EDGE, r=16, name="Concatenation bar")
tf = cat.text_frame
tf.word_wrap = False
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf._txBody.find(qn("a:bodyPr")).set("vert", "vert270")
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Concatenation: original + SR"
r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(FS_T), True, NAVY

# --- feature fusion stack
g = group("Block: Feature fusion stack")
sx, sy, sz = 900, 118, 240
stack = [f"{STEM}_input_rgb.png", f"{STEM}_ediffsr_rgb.png", f"{STEM}_input_ir.png", f"{STEM}_ediffsr_ir.png"]
names = ["Original RGB", "EDiffSR RGB", "Original IR composite", "EDiffSR IR composite"]
for k in (3, 2, 1, 0):
    pic(g, A / stack[k], sx + 26 * k, sy - 24 * k, sz, sz, name=f"Layer {k + 1}: {names[k]}")
label(slide.shapes, 1075, 418, 300, 100, "Feature fusion\noriginal + SR\n(16 bands)", name="Feature fusion label")
seg(slide.shapes, 848, 560, 876, 560, name="Line to stack 1")
seg(slide.shapes, 876, 560, 876, 250, name="Line to stack 2")
seg(slide.shapes, 876, 250, 898, 250, arrow=True, name="Arrow to stack")

# --- segmentation model
g = group("Block: Segmentation model")
rbox(g, 920, 510, 340, 130, SEG_BG, SEG_EDGE, name="Box")
label(g, 1090, 575, 320, 90, "Segmentation model\n(DeepLabV3+)", name="Title")
seg(slide.shapes, 1090, 484, 1090, 510, arrow=True, name="Arrow to model")
seg(slide.shapes, 1090, 640, 1090, 700, arrow=True, name="Arrow to map")

# --- map
g = group("Block: Cloud and shadow map")
rbox(g, 900, 700, 380, 380, MASK_BG, GREY_EDGE, r=14, name="Box")
pic(g, A / f"{STEM}_overlay.png", 925, 735, 330, 330, name="Picture: map overlay")
label(slide.shapes, 1090, 1112, 420, 36, "Cloud and cloud-shadow masks", FS_N, name="Map caption")

# --- legend
seg(slide.shapes, 1280, 840, 1330, 840, arrow=True, name="Arrow to legend")
g = group("Block: Legend")
rbox(g, 1330, 700, 240, 270, WHITE, GREY_EDGE, r=16, name="Box")
label(g, 1450, 735, 200, 36, "Classes", FS_N, name="Title")
for k, (nm, c) in enumerate([("Clear", CLEAR), ("Cloud", CLOUD), ("Shadow", SHADOW)]):
    y = 800 + k * 62
    sq = g.add_shape(MSO_SHAPE.RECTANGLE, E(1352), E(y - 18), E(36), E(36))
    sq.fill.solid()
    sq.fill.fore_color.rgb = c
    sq.line.color.rgb = NAVY
    sq.line.width = Pt(1)
    sq.shadow.inherit = False
    sq.name = f"Swatch {nm}"
    label(g, 0, y, 160, 36, nm, FS_N, bold=False, align=PP_ALIGN.LEFT, x_left=1405, name=f"Class {nm}")

# --- findings
g = group("Block: Findings")
rbox(g, 34, 1150, 1532, 150, WHITE, GREY_EDGE, r=20, name="Box")
items = [("SR restores edges better\nthan bicubic interpolation", 300),
         ("On cloudy patches, SR gives no\ngain over bicubic; IoU and recall\nof small and thin clouds rise\nslightly (not significant)", 750),
         ("Combining original data with SR in the\n16-channel feature space saves both\nspectrum and high quality details, needed\nfor climate structures alignment", 1280)]
for i, (t, x) in enumerate(items):
    label(g, x, 1225, 500, 120, t, FS_S + 1, name=f"Finding {i + 1}")
for x in (492, 1002):
    seg(g, x, 1170, x, 1280, color=RGBColor(0xC0, 0xC0, 0xC0), lw=1.5, name="Separator")

prs.save(OUT)
print("saved", OUT)
